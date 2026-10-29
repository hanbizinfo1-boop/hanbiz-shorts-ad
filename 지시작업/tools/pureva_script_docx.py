"""퓨레바 대본 3종 → 워드(한 쪽에 대본 하나) — 대본 글자는 손대지 않는다.

원본은 사장님이 한글로 고쳐 저장한 `퓨레바_대본_3종.hwp`(10-23 11:56). 그 안의 문단을 그대로 읽어
「① 육아」「② 부모님 선물」「③ 직장인」 제목 줄에서 나눈다. 쪽 넘김은 제목 문단 «쪽 나눔 앞»으로.
"""
import struct
import zlib
from pathlib import Path

import olefile
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

DIR = Path(__file__).resolve().parents[1] / "납품" / "퓨레바_대본_컨펌용_1023"
SRC = DIR / "퓨레바_대본_3종.hwp"
OUT = DIR / "퓨레바_대본_3종_한쪽씩.docx"
TITLES = ["① 육아", "② 부모님 선물", "③ 직장인"]

TAG_PARA_HEADER, TAG_PARA_TEXT = 66, 67
INLINE = {0, 10, 13, 24, 25, 26, 27, 28, 29, 30, 31}  # 2바이트 제어 문자 — 나머지는 16바이트


def hwp_paragraphs(path):
    ole = olefile.OleFileIO(str(path))
    data = ole.openstream("BodyText/Section0").read()
    if ole.openstream("FileHeader").read()[36] & 1:
        data = zlib.decompress(data, -15)
    paras, i = [], 0
    while i < len(data):
        head = struct.unpack_from("<I", data, i)[0]
        i += 4
        tag, level, size = head & 0x3FF, (head >> 10) & 0x3FF, head >> 20
        if size == 0xFFF:
            size = struct.unpack_from("<I", data, i)[0]
            i += 4
        if tag == TAG_PARA_HEADER and level == 0:
            paras.append("")
        elif tag == TAG_PARA_TEXT and level == 1:
            b, j, s = data[i:i + size], 0, []
            while j < len(b):
                c = struct.unpack_from("<H", b, j)[0]
                if c < 32:
                    j += 2 if c in INLINE else 16
                    continue
                s.append(chr(c))
                j += 2
            paras[-1] = "".join(s)
        i += size
    return paras


def set_font(run, size, bold=False):
    run.font.name = "맑은 고딕"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    return run


def split_scripts(paras):
    idx = [paras.index(t) for t in TITLES] + [len(paras)]
    out = []
    for k, t in enumerate(TITLES):
        body = paras[idx[k] + 1:idx[k + 1]]
        while body and body[-1] == "":
            body.pop()
        out.append((t, body))
    return out


scripts = split_scripts(hwp_paragraphs(SRC))

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(20)
sec.left_margin = sec.right_margin = Mm(25)
# 기본 서식의 문서 격자(줄 간격 18pt)에 맑은 고딕 줄이 걸리면 줄마다 격자 2칸을 먹어 한 대본이 두 쪽이 된다 — 격자를 뗀다
for grid in sec._sectPr.findall(qn("w:docGrid")):
    sec._sectPr.remove(grid)

for k, (title, body) in enumerate(scripts):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = k > 0
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run(title), 15, bold=True)
    for line in body:
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.line_spacing = 1.15
        set_font(q.add_run(line), 11)

doc.save(OUT)

# 검산 — 워드 속 대본 = hwp 문단, 글자 하나까지
paras = [p.text for p in Document(OUT).paragraphs]
for title, body in scripts:
    s = paras.index(title) + 1
    assert paras[s:s + len(body)] == body, title
    # 한 쪽 높이: 본문 257mm ≈ 728pt · 줄당 11pt×1.15×1.3(글꼴 여유) ≈ 16.4pt
    need = 15 * 1.3 + 10 + len(body) * 16.4
    print(f"{title}: {len(body)}줄 · 예상 {need:.0f}pt / 728pt")
print(OUT)
